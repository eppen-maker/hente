from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
SRC="PLACEHOLDER_DATA_ID"
NAVY="1F3864"; BAND="E8EFF7"; LINE="C9D3E0"
wb=Workbook(); ws=wb.active; ws.title="Dashbord"
ws.sheet_view.showGridLines=False

ws.merge_cells("A1:C1")
c=ws["A1"]; c.value="DASHBORD  -  TIMER +/-"
c.fill=PatternFill("solid",fgColor=NAVY)
c.font=Font(color="FFFFFF",bold=True,size=18)
c.alignment=Alignment(horizontal="left",vertical="center",indent=1)
ws.row_dimensions[1].height=32

c=ws["A2"]; c.value="Minustimer overst. Gronn = timer til gode, rod = timer som skyldes."
c.font=Font(color="1F3864",italic=True,size=10)
c.alignment=Alignment(horizontal="left",vertical="center")
ws.row_dimensions[2].height=18

for i,t in enumerate(["Navn","Timer +/-","Begrunnelse"]):
    c=ws.cell(row=4,column=1+i,value=t)
    c.fill=PatternFill("solid",fgColor=NAVY)
    c.font=Font(color="FFFFFF",bold=True,size=11)
    c.alignment=Alignment(horizontal="center",vertical="center")
ws.row_dimensions[4].height=28

ws["A5"]='=IMPORTRANGE("%s","A5:C17")'%SRC

t=Side(style="thin",color=LINE); bd=Border(left=t,right=t,top=t,bottom=t)
for r in range(5,18):
    ws.row_dimensions[r].height=28
    band = (r%2==0)
    for col in (1,2,3):
        x=ws.cell(row=r,column=col); x.border=bd
        if band or r==17: x.fill=PatternFill("solid",fgColor=BAND)
        if col==1: x.font=Font(bold=True); x.alignment=Alignment(vertical="center",indent=1)
        elif col==2:
            x.number_format="0.00"
            x.alignment=Alignment(horizontal="center",vertical="center")
        else: x.alignment=Alignment(wrap_text=True,vertical="center",indent=1)
ws.cell(row=17,column=2).font=Font(bold=True)

ws.conditional_formatting.add("B5:B17",CellIsRule(operator="greaterThan",formula=["0"],
    fill=PatternFill("solid",fgColor="D6F0DD"),font=Font(color="0B6B32",bold=True)))
ws.conditional_formatting.add("B5:B17",CellIsRule(operator="lessThan",formula=["0"],
    fill=PatternFill("solid",fgColor="FADBD8"),font=Font(color="A32218",bold=True)))

ws.column_dimensions["A"].width=21
ws.column_dimensions["B"].width=13
ws.column_dimensions["C"].width=92
ws.freeze_panes="A5"
wb.save("g8.xlsx")
